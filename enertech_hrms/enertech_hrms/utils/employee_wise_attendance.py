"""
Employee Wise Monthly Attendance Report (Excel download)

Place at:
  apps/enertech_hrms/enertech_hrms/enertech_hrms/utils/employee_wise_attendance.py

Whitelisted method:
  enertech_hrms.enertech_hrms.utils.employee_wise_attendance.download_report

Needs only openpyxl (ships with Frappe). No pandas.
"""
import calendar
from datetime import date
from io import BytesIO

import frappe
from frappe import _
from frappe.utils import cint, getdate
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

DAY_CODES = ["MON", "TUE", "WED", "THU", "FRI", "SAT", "SUN"]
ROW_LABELS = ["Shift", "In", "Out", "Break Hrs", "Net Hrs", "Extra Hrs",
              "Late Hrs", "Early Hrs", "Short Hrs", "Status", "Remark"]


# --------------------------------------------------------------------------
# helpers
# --------------------------------------------------------------------------
def _hhmm(minutes):
    minutes = int(round(minutes))
    return f"{minutes // 60:02d}:{minutes % 60:02d}"


def _hm(s):
    h, m = str(s).split(":")[:2]
    return int(h) * 60 + int(m)


def _time_to_min(dt):
    """datetime / timedelta / 'HH:MM:SS' string -> minutes from midnight."""
    if not dt:
        return None
    if hasattr(dt, "hour"):
        return dt.hour * 60 + dt.minute
    if hasattr(dt, "total_seconds"):  # timedelta
        return int(dt.total_seconds() // 60) % (24 * 60)
    s = str(dt)
    if " " in s:
        s = s.split(" ")[1]
    return _hm(s)


def _leave_code(leave_type):
    if not leave_type:
        return "L"
    s = leave_type.strip().lower()
    known = {"casual leave": "CL", "sick leave": "SL", "privilege leave": "PL",
             "earned leave": "EL", "compensatory off": "CO", "comp off": "CO"}
    return known.get(s) or "".join(w[0] for w in s.split()).upper()


def _fmt_num(v):
    return int(v) if float(v).is_integer() else v


# --------------------------------------------------------------------------
# data access
# --------------------------------------------------------------------------
def _get_holiday_map(employee, from_date, to_date, cache):
    """{date: is_weekly_off}  from the employee's Holiday List (cached per list)."""
    hl = None
    try:
        from erpnext.setup.doctype.employee.employee import get_holiday_list_for_employee
        hl = get_holiday_list_for_employee(employee, raise_exception=False)
    except Exception:
        pass

    if not hl:
        return None  # caller falls back to Sunday

    if hl not in cache:
        rows = frappe.get_all(
            "Holiday",
            filters={"parent": hl, "holiday_date": ["between", [from_date, to_date]]},
            fields=["holiday_date", "weekly_off"],
        )
        cache[hl] = {getdate(r.holiday_date): bool(r.weekly_off) for r in rows}
    return cache[hl]


def _get_data(year, month, company=None):
    ndays = calendar.monthrange(year, month)[1]
    from_date, to_date = date(year, month, 1), date(year, month, ndays)

    att_filters = {"attendance_date": ["between", [from_date, to_date]], "docstatus": 1}
    emp_filters = {"status": "Active"}
    if company:
        att_filters["company"] = company
        emp_filters["company"] = company

    attendance = frappe.get_all(
        "Attendance",
        filters=att_filters,
        fields=["employee", "employee_name", "status", "attendance_date",
                "in_time", "out_time", "working_hours", "leave_type"],
        order_by="attendance_date asc",
        limit_page_length=0,
    )

    # active employees (so people with zero records still show) + anyone who has records
    employees = {e.name: e.employee_name for e in frappe.get_all(
        "Employee", filters=emp_filters, fields=["name", "employee_name"], limit_page_length=0)}
    for a in attendance:
        employees.setdefault(a.employee, a.employee_name)

    by_emp = {}
    for a in attendance:
        by_emp.setdefault(a.employee, {}).setdefault(getdate(a.attendance_date).day, a)

    return employees, by_emp, from_date, to_date, ndays


# --------------------------------------------------------------------------
# report builder
# --------------------------------------------------------------------------
def build_workbook(employees, by_emp, holiday_map_for, year, month, ndays,
                   location, shift_start, shift_end, grace):
    s_start, s_end = _hm(shift_start), _hm(shift_end)
    shift_len = s_end - s_start
    month_label = date(year, month, 1).strftime("%b-%Y")

    F = "Arial"
    f_norm = Font(name=F, size=9)
    f_bold = Font(name=F, size=9, bold=True)
    f_red = Font(name=F, size=9, color="FF0000")
    f_title = Font(name=F, size=11, bold=True)
    green = PatternFill("solid", fgColor="D8E4BC")
    thin = Side(style="thin", color="000000")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center")
    left = Alignment(horizontal="left", vertical="center")

    wb = Workbook()
    ws = wb.active
    ws.title = (location or "Attendance")[:31]

    FIRST = 3
    SUMCOL = FIRST + ndays
    ws.cell(1, 1, f"Employee Wise Attendance Report - ({month_label}) - {location}").font = f_title
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 11
    for i in range(ndays):
        ws.column_dimensions[get_column_letter(FIRST + i)].width = 7.2
    ws.column_dimensions[get_column_letter(SUMCOL)].width = 24

    row = 3
    for emp_id, emp_name in sorted(employees.items(), key=lambda kv: (kv[1] or "").upper()):
        recs = by_emp.get(emp_id, {})
        hmap = holiday_map_for(emp_id)  # {date: is_weekly_off} or None -> Sunday fallback
        r0 = row

        # header row
        ws.cell(r0, 1, "EmployeeName")
        ws.cell(r0, 2, "Date")
        for d in range(1, ndays + 1):
            c = ws.cell(r0, FIRST + d - 1, date(year, month, d).strftime("%d-%m-%y"))
            c.alignment = center
        ws.cell(r0, SUMCOL, "Summary")
        for c in range(1, SUMCOL + 1):
            cell = ws.cell(r0, c)
            cell.font, cell.fill, cell.border = f_bold, green, box

        for i, lbl in enumerate(ROW_LABELS, start=1):
            ws.cell(r0 + i, 1).border = box
            ws.cell(r0 + i, 2, lbl).font = f_norm
            ws.cell(r0 + i, 2).border = box
        ws.cell(r0 + 1, 1, f"{emp_name} ({emp_id})").font = f_bold

        tot = dict(P=0.0, AB=0.0, WO=0, LV=0.0, NET=0.0, OT=0.0, LATE=0.0, EARLY=0.0, SHORT=0.0, PH=0)

        for d in range(1, ndays + 1):
            dt = date(year, month, d)
            col = FIRST + d - 1

            if hmap is None:
                is_wo, is_ph = dt.weekday() == 6, False
            elif dt in hmap:
                is_wo, is_ph = hmap[dt], not hmap[dt]
            else:
                is_wo = is_ph = False

            rec = recs.get(d)
            st = (rec.status if rec else "") or ""
            vals = {k: "" for k in ROW_LABELS}
            vals["Shift"] = f"{shift_start} - {shift_end}"
            for k in ("Break Hrs", "Net Hrs", "Extra Hrs", "Late Hrs", "Early Hrs", "Short Hrs"):
                vals[k] = "00:00"
            status = ""

            if rec and st in ("Present", "Work From Home", "Half Day"):
                tin, tout = _time_to_min(rec.in_time), _time_to_min(rec.out_time)
                wh = float(rec.working_hours or 0) * 60
                if not wh and tin is not None and tout is not None and tout > tin:
                    wh = tout - tin
                if tin is not None:
                    vals["In"] = _hhmm(tin)
                if tout is not None:
                    vals["Out"] = _hhmm(tout)
                vals["Net Hrs"] = _hhmm(wh)
                tot["NET"] += wh
                if st != "Half Day":
                    if wh > shift_len:
                        vals["Extra Hrs"] = _hhmm(wh - shift_len); tot["OT"] += wh - shift_len
                    if tin is not None and tin > s_start + grace:
                        vals["Late Hrs"] = _hhmm(tin - s_start); tot["LATE"] += tin - s_start
                    if tout is not None and tout < s_end:
                        vals["Early Hrs"] = _hhmm(s_end - tout); tot["EARLY"] += s_end - tout
                    if 0 < wh < shift_len:
                        vals["Short Hrs"] = _hhmm(shift_len - wh); tot["SHORT"] += shift_len - wh

            if rec and st in ("Present", "Work From Home"):
                status = "WFH" if st == "Work From Home" else "P"
                tot["P"] += 1
                if is_wo:
                    status = "WO(P)"; tot["WO"] += 1
                elif is_ph:
                    status = "PH(P)"; tot["PH"] += 1
            elif rec and st == "Half Day":
                other = _leave_code(rec.leave_type) if rec.leave_type else "AB"
                status = f"P/{other}"
                tot["P"] += 0.5
                tot["AB" if other == "AB" else "LV"] += 0.5
            elif rec and st == "On Leave":
                if not rec.leave_type or "without pay" in rec.leave_type.lower():
                    status = "AB"; tot["AB"] += 1
                else:
                    status = _leave_code(rec.leave_type); tot["LV"] += 1
            elif rec and st == "Absent":
                if is_wo:
                    status = "WO"; tot["WO"] += 1
                elif is_ph:
                    status = "PH"; tot["PH"] += 1
                else:
                    status = "AB"; tot["AB"] += 1
            else:  # no record
                if is_ph:
                    status = "PH"; tot["PH"] += 1
                elif is_wo:
                    status = "WO"; tot["WO"] += 1

            vals["Status"] = status
            for i, lbl in enumerate(ROW_LABELS, start=1):
                c = ws.cell(r0 + i, col, vals[lbl])
                c.font = f_red if (lbl == "Status" and "AB" in status) else f_norm
                c.alignment, c.border = center, box

        # Remark row (green)
        for c in range(1, SUMCOL + 1):
            ws.cell(r0 + 11, c).fill = green
            ws.cell(r0 + 11, c).border = box
        ws.cell(r0 + 11, 2, "Remark").font = f_norm

        summary = [
            f"Present : {_fmt_num(tot['P'])}",
            f"Absent : {_fmt_num(tot['AB'])}",
            f"Weekly Off : {tot['WO']}",
            f"Leaves : {_fmt_num(tot['LV'])}",
            f"Net Hrs : {tot['NET'] / 60:.2f}",
            f"OT Hrs : {tot['OT'] / 60:.2f}",
            f"Late In Hrs : {tot['LATE'] / 60:.2f}",
            f"Early Out Hrs : {tot['EARLY'] / 60:.2f}",
            f"Short Hrs : {tot['SHORT'] / 60:.2f}",
            f"Holiday : {tot['PH']}",
            "CO Days : 0",
        ]
        for i, txt in enumerate(summary, start=1):
            c = ws.cell(r0 + i, SUMCOL, txt)
            c.font, c.border, c.alignment = f_norm, box, left

        row = r0 + 12

    ws.freeze_panes = "C3"
    ws.sheet_view.zoomScale = 90
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue(), month_label


# --------------------------------------------------------------------------
# whitelisted entry point
# --------------------------------------------------------------------------
@frappe.whitelist()
def download_report(month=None, year=None, location="Pirangut", company=None,
                    shift_start="09:30", shift_end="18:00", grace=15):
    """Download Employee Wise Attendance Report as .xlsx

    /api/method/enertech_hrms.enertech_hrms.utils.employee_wise_attendance.download_report?month=9&year=2026
    """
    frappe.only_for(["HR Manager", "HR User", "System Manager"])

    today = getdate()
    month = cint(month) or today.month
    year = cint(year) or today.year
    if not 1 <= month <= 12:
        frappe.throw(_("Invalid month"))

    employees, by_emp, from_date, to_date, ndays = _get_data(year, month, company)
    if not employees:
        frappe.throw(_("No employees / attendance found for the selected month"))

    cache = {}
    emp_cache = {}

    def holiday_map_for(emp_id):
        if emp_id not in emp_cache:
            emp_cache[emp_id] = _get_holiday_map(emp_id, from_date, to_date, cache)
        return emp_cache[emp_id]

    content, month_label = build_workbook(
        employees, by_emp, holiday_map_for, year, month, ndays,
        location, shift_start, shift_end, cint(grace),
    )

    frappe.response["filename"] = f"Employee_Wise_Attendance_{month_label}.xlsx"
    frappe.response["filecontent"] = content
    frappe.response["type"] = "binary"